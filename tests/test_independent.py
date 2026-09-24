# SPDX-License-Identifier: AGPL-3.0-only
# JPM modified distribution, 2026-09-24. See package-root NOTICE.md.
"""Behavior checks for the independent fragment inspector and optional Word reader."""
from pathlib import Path
import json
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from check_fragment import inspect_fragment

class FragmentTests(unittest.TestCase):
    def test_pasteable_content(self):
        html = '<section style="background:linear-gradient(90deg,#123,#456);"><p><span leaf="">中文 &amp; punctuation: unchanged.</span></p><a href="https://example.com/path?q=1&amp;b=2"><span leaf="">来源</span></a></section>'
        result = inspect_fragment(html)
        self.assertTrue(result['passed'], result)
        self.assertEqual(result['counts']['links'], 1)

    def test_executable_or_layout_incompatible_content_is_rejected(self):
        cases = [
            '<section><script>alert(1)</script></section>',
            '<section onclick="alert(1)"><span leaf="">X</span></section>',
            '<section><a href="java&#x73;cript:alert(1)"><span leaf="">X</span></a></section>',
            '<section style="background:u\\72l(https://example.com/)"><span leaf="">X</span></section>',
            '<section style="position: fixed !important"><span leaf="">X</span></section>',
            '<section style="display:inline-grid"><span leaf="">X</span></section>',
            '<section class="card"><span leaf="">X</span></section>',
            '<section>Visible unwrapped text</section>',
            '<section><p><section><span leaf="">X</span></section></p></section>',
            '<section><p><span leaf="">X</p></span></section>',
            '<section></section><section></section>',
            '<!DOCTYPE html><section></section>',
        ]
        for html in cases:
            with self.subTest(html=html):
                self.assertFalse(inspect_fragment(html)['passed'])

    def test_relative_image_exists_and_does_not_escape(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'image.png').write_bytes(b'fixture')
            def fragment(src):
                return '<section><span leaf=""><img src="'+src+'" style="max-width:100%;height:auto;"></span></section>'
            self.assertTrue(inspect_fragment(fragment('image.png'), root)['passed'])
            self.assertFalse(inspect_fragment(fragment('missing.png'), root)['passed'])
            self.assertFalse(inspect_fragment(fragment('../image.png'), root)['passed'])
            self.assertFalse(inspect_fragment(fragment('data:image/png;base64,AA'), root)['passed'])

class WordTests(unittest.TestCase):
    def test_order_images_links_tables_and_unsupported_content(self):
        try:
            from docx import Document
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
        except ImportError:
            self.skipTest('Optional python-docx dependency is absent')
        from read_word import collect_document
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            document = Document()
            document.add_heading('客户观察记录', 0)
            p = document.add_paragraph()
            p.add_run('图前内容')
            run = p.add_run('强调'); run.bold = True; run.underline = True
            p.add_run().add_picture(str(ROOT / 'evals/fixtures/layout-dot.png'))
            p.add_run('图后内容')
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), p.part.relate_to('https://example.com/source', RT.HYPERLINK, is_external=True))
            r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '资料来源'; r.append(t); link.append(r); p._p.append(link)
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = '合并标题'
            table.cell(1, 0).text = '次数'; table.cell(1, 1).text = '12'
            paragraph = document.add_paragraph('结束')
            insertion = OxmlElement('w:ins'); paragraph._p.append(insertion)
            source = root / 'sample.docx'; document.save(source)
            output = root / 'review'; report = collect_document(source, output)
            tokens = report['records'][1]['tokens']
            kinds = [t['kind'] for t in tokens]
            self.assertEqual(kinds, ['text','text','image','text','link'])
            self.assertEqual(tokens[3]['value'], '图后内容')
            self.assertTrue(tokens[1]['bold'] and tokens[1]['underline'])
            self.assertEqual(tokens[-1]['url'], 'https://example.com/source')
            self.assertEqual((output / tokens[2]['source']).read_bytes(), (ROOT/'evals/fixtures/layout-dot.png').read_bytes())
            grid = report['records'][2]
            self.assertEqual(grid['rows'][0]['cells'], [0, 0])
            self.assertEqual(len(grid['cells']), 3)
            self.assertTrue(any('tracked insertions' in w for w in report['warnings']))
            self.assertTrue(report['reviewRequired'])
            self.assertTrue((output/'word-review.json').is_file())
            with self.assertRaises(ValueError): collect_document(source, output)

if __name__ == '__main__':
    unittest.main()
